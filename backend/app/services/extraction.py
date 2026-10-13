"""Plain-text extraction for PDF (pypdf) and DOCX (python-docx).

Every function returns an ExtractionResult instead of raising for
expected document problems (empty, encrypted, image-only, corrupt),
so the API can persist an honest status. Nothing here contacts any
external service — text stays on this server (Phase 2 will reuse it
for chunking/RAG).
"""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

from docx import Document as DocxDocument
from pypdf import PdfReader
from pypdf.errors import PdfReadError


@dataclass
class ExtractionResult:
    text: str | None
    status: str  # completed | failed | requires_ocr
    error: str | None = None
    page_count: int | None = None


def extract_pdf_text(path: str | Path) -> ExtractionResult:
    try:
        reader = PdfReader(str(path))
    except (PdfReadError, OSError, ValueError) as exc:
        return ExtractionResult(
            text=None, status="failed", error=f"Unreadable PDF: {exc}"
        )

    if reader.is_encrypted:
        try:
            # Some producers "encrypt" with a blank password; try that once.
            result = reader.decrypt("")
            if result == 0:
                return ExtractionResult(
                    text=None,
                    status="failed",
                    error="This PDF is password-protected.",
                )
        except Exception:
            return ExtractionResult(
                text=None,
                status="failed",
                error="This PDF is password-protected.",
            )

    page_count = len(reader.pages)
    if page_count == 0:
        return ExtractionResult(
            text=None, status="failed", error="This PDF has no pages.", page_count=0
        )

    pages: list[str] = []
    for index, page in enumerate(reader.pages, start=1):
        try:
            text = page.extract_text() or ""
        except Exception:
            text = ""
        cleaned = text.strip()
        if cleaned:
            pages.append(f"--- Page {index} ---\n{cleaned}")

    if not pages:
        return ExtractionResult(
            text=None,
            status="requires_ocr",
            error="No selectable text found — this looks like a scanned document.",
            page_count=page_count,
        )
    return ExtractionResult(
        text="\n\n".join(pages), status="completed", page_count=page_count
    )


def extract_docx_text(path: str | Path) -> ExtractionResult:
    try:
        document = DocxDocument(str(path))
    except Exception as exc:
        return ExtractionResult(
            text=None, status="failed", error=f"Unreadable DOCX: {exc}"
        )

    parts = [p.text.strip() for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))

    text = "\n\n".join(part for part in parts if part).strip()
    if not text:
        return ExtractionResult(
            text=None, status="failed", error="This DOCX contains no text."
        )
    return ExtractionResult(text=text, status="completed")
